"""Lees de eigen conversievoorbeelden onafhankelijk terug; geen volledige opmaakcontrole."""
from pathlib import Path
import json
import sys
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
uit=Path(sys.argv[1])
bron=Path(__file__).resolve().parent/'voorbeelden'
metingen=json.loads((uit/'metingen.json').read_text())
assert len(metingen)==15

def tekst(p):
    d=Document(p)
    return {'alineas':[a.text for a in d.paragraphs], 'tabellen':[[[c.text for c in r.cells] for r in t.rows] for t in d.tables]}
def cellen(p):
    w=load_workbook(p,data_only=False)
    return {s.title:{c.coordinate:c.value for r in s for c in r if c.value is not None} for s in w}
def dias(p):
    def vormen(s):
        gevonden=[]
        for v in s:
            if v.has_text_frame: gevonden.append(v.text)
            if hasattr(v,'shapes'): gevonden.extend(vormen(v.shapes))
        return gevonden
    return [vormen(s.shapes) for s in Presentation(p).slides]
resultaten=[]
for naam,ext,oud,odf,lees in [('Ontwerpnotitie','docx','doc','odt',tekst),('Materiaaloverzicht','xlsx','xls','ods',cellen),('Ontwerpvisie','pptx','ppt','odp',dias)]:
    verwacht=lees(bron/f'{naam}.{ext}')
    for via in [oud,odf]:
        pad=uit/f'{naam}-via-{via}.{ext}'
        assert lees(pad)==verwacht,f'Inhoud wijkt af: {pad.name}'
        assert (uit/f'{naam}.{via}').stat().st_size>100
        resultaten.append({'bestand':pad.name,'inhoud_gelijk':True})
    assert (uit/f'{naam}.pdf').read_bytes().startswith(b'%PDF-')
print(json.dumps({'rondgangen':resultaten,'metingen':len(metingen),'pdf_handtekeningen':3,'omvang':'Alineatekst/tabellen, niet-lege celinhoud/formules en diatekst; geen volledige stijl- of pixelcontrole'},ensure_ascii=False))
