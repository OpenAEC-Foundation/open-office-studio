"""Eigen driedelige paginaproef voor eerste, even en oneven kop/voet."""
from pathlib import Path
from docx import Document
map=Path(__file__).parent
d=Document()
d.settings.odd_and_even_pages_header_footer=True
s=d.sections[0]
s.different_first_page_header_footer=True
# De vaste aanmaakvolgorde bepaalt header/footer1, 2 en 3 voor de proef.
s.header.paragraphs[0].text='Oneven koptekst'
s.even_page_header.paragraphs[0].text='Even koptekst'
s.first_page_header.paragraphs[0].text='Eerste koptekst'
s.footer.paragraphs[0].text='Oneven voettekst'
s.even_page_footer.paragraphs[0].text='Even voettekst'
s.first_page_footer.paragraphs[0].text='Eerste voettekst'
d.add_paragraph('Eerste blad')
d.add_page_break()
d.add_paragraph('Tweede blad')
d.add_page_break()
d.add_paragraph('Derde blad')
d.save(map/'writer-kopvoet.docx')
