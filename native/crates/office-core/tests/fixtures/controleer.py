"""Lees native uitvoer met drie onafhankelijke documentbibliotheken."""
import sys
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
map = Path(sys.argv[1])
doc = Document(map/'writer.docx')
assert any(p.text == 'Nieuwe tekst met café en 🏠' for p in doc.paragraphs)
assert doc.tables[0].cell(1,0).text == 'Tabel inhoud'
assert doc.paragraphs[2].runs[0].bold
assert doc.paragraphs[2].runs[1].italic
rijk = Document(map/'writer-rijk.docx')
assert rijk.paragraphs[2].text == 'Behoud deze cursieve tekst.\nTweede regel met 🏠'
assert rijk.paragraphs[2].runs[0].bold
assert any(r.italic for r in rijk.paragraphs[2].runs if 'cursieve' in r.text)
wb = load_workbook(map/'spreadsheet.xlsx', data_only=False)
assert wb['Begroting']['B2'].value == 4
assert wb['Begroting']['D2'].value == '=B2*C2'
assert wb['Begroting']['D4'].value == '=SUM(D2:D3)'
assert wb['Begroting']['D4'].number_format == '0.00'
assert wb['Begroting']['A1'].font.bold
assert wb['Tweede blad']['A1'].value == 'Bewaar ook dit blad'
waarden = load_workbook(map/'spreadsheet.xlsx', data_only=True)
assert waarden['Begroting']['D2'].value == 50
assert waarden['Begroting']['D4'].value == 80
ppt = Presentation(map/'presentation.pptx')
assert len(ppt.slides) == 2
assert any(s.has_text_frame and s.text == 'Gewijzigde diatekst' for s in ppt.slides[0].shapes)
assert 'Sprekersnotities' in ppt.slides[0].notes_slide.notes_text_frame.text
assert any(s.has_text_frame and s.text == 'Tweede dia blijft bewaard' for s in ppt.slides[1].shapes)
rijk = Presentation(map/'presentation-rijk.pptx')
p = rijk.slides[0].shapes[0].text_frame.paragraphs[0]
assert p.text == 'Native laatste 🏠 tekst'
assert p.runs[0].text == 'Native ' and p.runs[0].font.bold
print('Native DOCX/XLSX/PPTX: inhoud, stijlen, formules, rekenwaarden, tabellen, notities en tweede blad/dia onafhankelijk goedgekeurd.')

opmaak = Document(map/'writer-opmaak.docx')
p = next(p for p in opmaak.paragraphs if p.text == 'Behoud deze cursieve tekst.')
assert p.runs[0].text == 'Behoud ' and p.runs[0].bold
assert next(r for r in p.runs if r.text == 'deze').underline
assert next(r for r in p.runs if r.text == 'deze').italic
assert int(p.alignment) == 1
letter = Document(map/'writer-lettergrootte.docx')
p = next(p for p in letter.paragraphs if p.text == 'Bewerkbare oorspronkelijke alinea.')
assert p.runs[0].font.size.pt == 10.5
structuur = load_workbook(map/'spreadsheet-structuur.xlsx', data_only=False)
s = structuur['Begroting']
assert s['D3'].value == '=B3*C3'
assert s['B3'].font.bold and s['C4'].font.bold
assert s['B3'].fill.fgColor.rgb.upper().endswith('FFF2CC')
assert structuur['Tweede blad']['A1'].value == 'Bewaar ook dit blad'
print('Native uitbreidingen: tekstbereikopmaak, halve punten en celstijl/rij-invoeging onafhankelijk goedgekeurd.')

dubbel = Document(map/'writer-dubbele-ids.docx')
assert dubbel.paragraphs[0].text == 'Native documentcontrole'
assert dubbel.paragraphs[1].text == 'Alleen deze alinea wijzigt.'
print('Herhaalde alinea-identiteiten: uitsluitend de bedoelde alinea is veranderd.')

gesplitst = Document(map/'writer-gesplitst.docx')
assert [p.text for p in gesplitst.paragraphs][:3] == ['Voor de rijke alinea', 'Links 🏠 ', 'café rechts']
assert gesplitst.paragraphs[1].runs[0].bold
assert gesplitst.paragraphs[2].runs[0].italic
assert len(gesplitst.sections) == 2
samengevoegd = Document(map/'writer-samengevoegd.docx')
assert samengevoegd.paragraphs[1].text == 'Links 🏠 café rechts'
assert samengevoegd.paragraphs[1].runs[0].bold
assert samengevoegd.paragraphs[1].runs[-1].italic
assert len(samengevoegd.sections) == 2
assert samengevoegd.sections[0].header.paragraphs[0].text == 'Vaste koptekst'
assert samengevoegd.sections[0].footer.paragraphs[0].text == 'Vaste voettekst'
cel = Document(map/'writer-celstructuur.docx').tables[0]
assert [p.text for p in cel.cell(0,0).paragraphs] == ['Cel één', 'Tweede celalinea']
assert cel.cell(0,1).text == 'Cel twee'
print('Writer-alineastructuur: echte alinea’s, Unicode, runs, secties, kop/voettekst en celgrenzen onafhankelijk goedgekeurd.')

anker=load_workbook(map/'spreadsheet-anker.xlsx',data_only=False).active
assert anker['A1'].value == 'Gewijzigd via gedekt adres'
assert anker['H100'].value is None
assert 'A1:Z150' in [str(r) for r in anker.merged_cells.ranges]
korter=Document(map/'writer-pagina-weg.docx')
assert len(korter.paragraphs)==2
assert 'Bewerkbare oorspronkelijke alinea.' in korter.paragraphs[0].text
assert len(korter.tables)==1
print('Celanker buiten beeld en verdwijnen van de laatste Writer-pagina onafhankelijk teruggelezen.')

from zipfile import ZipFile
import xml.etree.ElementTree as ET
noot=Document(map/'writer-voetnoot.docx')
assert noot.paragraphs[0].text=='Tekst met een voetnoot.'
with ZipFile(map/'writer-voetnoot.docx') as z:
    wortel=ET.fromstring(z.read('word/footnotes.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    n=wortel.find('w:footnote',ns)
    assert n.attrib['{'+ns['w']+'}id']=='1'
    assert ''.join(n.itertext())=='Natieve voetnoot'
with ZipFile(map/'writer-grafiek-lettertype.docx') as z:
    xml=z.read('word/charts/chart1.xml')
    assert b'typeface="Gelasio"' in xml
    assert b'<c:v>2</c:v>' in xml and b'<c:v>4</c:v>' in xml
print('Voetnootnummer, noottekst en grafiekbron met eigen lettertype onafhankelijk teruggelezen.')

with ZipFile(map/'writer-velden.docx') as z:
    wortel=ET.fromstring(z.read('word/document.xml'))
    velden=wortel.findall('.//w:fldSimple',ns)
    assert len(velden)==2
    assert velden[0].find('w:fldSimple',ns) is not None
    assert ''.join(t.text or '' for t in velden[0].iter('{'+ns['w']+'}t'))=='Verwijzing 7'
    teksten=[''.join(t.text or '' for t in p.iter('{'+ns['w']+'}t')) for p in wortel.findall('w:body/w:p',ns)]
    assert teksten==['Inhoud café 42 🏠','Einde inhoud','Verwijzing 7']
    assert any('PAGEREF hoofdstuk' in (t.text or '') for t in wortel.iter('{'+ns['w']+'}instrText'))
print('Geneste velden: Unicode-resultaten, instructies en geneste veldstructuur onafhankelijk teruggelezen.')

# Pagina-instelling verandert uitsluitend sectie 1. De kop/voettekst van
# sectie 2 blijft geërfd; die verwijzingen mogen niet opnieuw worden ingebakken.
pagina = Document(map/'writer-pagina.docx')
assert len(pagina.sections) == 2
assert pagina.sections[0].page_width.twips == 16838
assert pagina.sections[0].page_height.twips == 11906
assert pagina.sections[1].page_width.twips == 12240
assert pagina.sections[1].page_height.twips == 15840
assert not pagina.sections[0].header.is_linked_to_previous
assert pagina.sections[1].header.is_linked_to_previous
assert pagina.sections[0].header.paragraphs[0].text == 'Vaste koptekst'
assert pagina.sections[1].footer.paragraphs[0].text == 'Vaste voettekst'
print('Writer-pagina-instelling: papier, stand en sectiegebonden kop-/voettekst onafhankelijk teruggelezen.')

# Stijlketen en de opdracht uit een echt bediende lintknop.
erfenis = Document(map/'writer-stijlerfenis.docx')
assert [p.text for p in erfenis.paragraphs] == ['Overgeërfde kop', 'Tekenstijl keert vet om', 'Directe opmaak']
assert erfenis.paragraphs[0].style.style_id == 'EigenKop'
assert erfenis.paragraphs[0].runs[0].bold is False
assert erfenis.styles['Eigen kop'].base_style.style_id == 'EigenBasis'
assert erfenis.styles['Eigen basis'].font.bold is True
assert erfenis.styles['Eigen kop'].font.size.pt == 18
assert erfenis.paragraphs[1].runs[0].style.style_id == 'EigenNadruk'
assert erfenis.paragraphs[2].runs[0].font.size.pt == 10.5
with ZipFile(map/'writer-stijlerfenis-bron.docx') as origineel, ZipFile(map/'writer-stijlerfenis.docx') as bewerkt:
    assert origineel.read('word/styles.xml') == bewerkt.read('word/styles.xml')
selectie = Document(map/'writer-selectie-lint.docx')
p = next(p for p in selectie.paragraphs if p.text == 'Links 🏠 café rechts')
assert p.runs[0].text == 'Links 🏠 ' and p.runs[0].bold is False
assert p.runs[-1].italic is True
print('Overgeërfde stijlen en Unicode-selectie via de echte lintknop onafhankelijk teruggelezen.')
for nummer, grootte in enumerate([10, 10.5, 10.5]):
    doc = Document(map/f'writer-letterprecisie-{nummer}.docx')
    p = next(p for p in doc.paragraphs if p.text == 'Bewerkbare oorspronkelijke alinea.')
    assert p.runs[0].font.size.pt == grootte
print('Afgeronde halve punten komen overeen met de weergegeven fontgrootte.')
