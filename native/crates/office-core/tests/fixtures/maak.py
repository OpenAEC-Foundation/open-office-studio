"""Synthetische bronbestanden voor onafhankelijke native rondgangcontroles."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
map = Path(__file__).parent
writer = Document()
writer.add_heading('Native documentcontrole', 0)
writer.add_paragraph('Bewerkbare oorspronkelijke alinea.')
p = writer.add_paragraph()
p.add_run('Behoud ').bold = True
p.add_run('deze cursieve tekst.').italic = True
t = writer.add_table(rows=2, cols=2)
for c, tekst in zip([c for r in t.rows for c in r.cells], ['Tabel kop', 'Waarde', 'Tabel inhoud', '42']): c.text = tekst
writer.save(map/'writer.docx')
tabel = Document()
t = tabel.add_table(rows=1, cols=1)
t.cell(0,0).text = 'Vier woorden in tabel'
tabel.save(map/'tabel.docx')
boek = Workbook()
s = boek.active
s.title = 'Begroting'
for r in [('Omschrijving', 'Aantal', 'Prijs', 'Totaal'), ('Onderdeel A', 3, 12.5, '=B2*C2'), ('Onderdeel B', 2, 15, '=B3*C3'), ('Totaal', None, None, '=SUM(D2:D3)')]: s.append(r)
s['D4'].number_format = '0.00'
s['A1'].font = s['A1'].font.copy(bold=True)
boek.create_sheet('Tweede blad')['A1'] = 'Bewaar ook dit blad'
boek.save(map/'spreadsheet.xlsx')
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
box.text = 'Native presentatiecontrole'
slide.shapes.add_textbox(Inches(1), Inches(3), Inches(7), Inches(1)).text = 'Bewerk deze tekst'
slide2 = deck.slides.add_slide(deck.slide_layouts[6])
slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = 'Tweede dia blijft bewaard'
deck.save(map/'presentation.pptx')
rijk = Presentation()
s = rijk.slides.add_slide(rijk.slide_layouts[6])
p = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.paragraphs[0]
r = p.add_run(); r.text = 'Native '; r.font.bold = True
r = p.add_run(); r.text = 'presentatiecontrole'; r.font.italic = True
rijk.save(map/'presentation-rijk.pptx')
